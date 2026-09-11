"""模板填数导出测试
==================================================================
合成 research run 与 reduction run，导出后读回校验：
  - sheet 名/各节标题/列头与模板逐字一致；
  - 橙色示例行与"↑以上为格式示例"说明行已删除；
  - 数据行数=输入条数；reduction run 含 Section 2/3；
  - Section 6 年行数=(horizon-baseline+1)×情景×子类别。
"""

import io
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "webapp"))

from export_excel import export_run_to_xlsx, TEMPLATE_PATH

CITE = {"source_type": "other", "title": "测试来源", "publisher": "测试出版方",
        "url": "https://example.org/x", "confidence": "low"}


def _research_result():
    return {
        "product_id": "pv_station", "product_name": "光伏电站",
        "region_id": "China", "region_name": "中国",
        "research_version": 1, "research_status": "approved",
        "classification_dimensions": [{
            "dimension_id": "dim-1", "dimension_name": "电池技术路线",
            "dimension_category": "technology_route",
            "axis_of_variation": "不同的电池片技术路线",
            "definition_rationale": "r", "orthogonality_check": "o",
            "values": [
                {"value_id": "v1", "value_name": "TOPCon晶硅组件",
                 "axis_conformity_justification": "一种电池片技术",
                 "market_share": {"value": 87.6, "unit": "%"},
                 "market_share_caliber": "中国电池技术路线市场占比",
                 "market_share_year": 2025, "region_basis": "region_specific",
                 "citations": [CITE]},
                {"value_id": "v2", "value_name": "BC组件",
                 "axis_conformity_justification": "一种电池片技术",
                 "region_basis": "global_fallback",
                 "region_basis_note": "区域证据不足，采用全球主流型号替代",
                 "citations": [CITE]},
            ],
            "evidence": [CITE],
        }],
        "functional_subsystems": [{
            "subsystem_id": "sub-1", "subsystem_name": "组件本体",
            "decomposition_level": 0, "function_description": "f", "evidence": [CITE]}],
        "copper_components": [{
            "component_id": "cu1", "parent_subsystem_id": "sub-1",
            "component_name": "互联焊带", "copper_form": "copper_foil",
            "function_of_copper": "导电", "mass_data_basis": "teardown_measurement",
            "unit_mass": {"value": 0.3559, "unit": "t"}, "citations": [CITE]}],
        "validation_flags": [], "objections": [], "debate_log": [],
    }


def _reduction_result():
    baselines = [{
        "value_id": "v1", "value_name": "TOPCon晶硅组件", "applicable_scope": None,
        "functional_unit": "t Cu/MWp_DC",
        "intensity": {"value": 0.5733, "unit": "t Cu/MWp_DC"},
        "component_weights": [
            {"component_id": "cu1", "component_name": "互联焊带",
             "mass": 0.3559, "mass_unit": "t", "weight_pct": 62.08},
            {"component_id": "cu2", "component_name": "汇流条",
             "mass": 0.2174, "mass_unit": "t", "weight_pct": 37.92},
        ],
        "is_reference": True, "high_copper_reference_note": "铜互联参考配置",
        "region_basis": "region_specific", "calc_trace": "test", "citations": [CITE],
    }]
    scenarios = []
    deltas = {"S0": 0.0, "S1": 3.74, "S2": 6.9, "S3": 9.5}
    names = {"S0": "基准情景", "S1": "普通减量", "S2": "加速减量", "S3": "深度减量"}
    for sid, d in deltas.items():
        scenarios.append({
            "scenario_id": sid, "scenario_name": names[sid], "applicable_scope": None,
            "scenario_definition": f"{sid}定义", "included_measure_ids": [] if sid == "S0" else ["R01"],
            "full_implementation_reduction_pct": {"value": d, "unit": "%"},
            "measure_start_year": 2025, "target_achievement_year": 2035,
            "target_achievement_rate_pct": 100, "diffusion_method": "linear",
            "scenario_rationale": "r", "citations": [CITE],
        })
    trajectory = []
    for sid, d in deltas.items():
        for year in range(2025, 2036):
            rate = 0.0 if year <= 2025 else (100.0 if year >= 2035 else (year - 2025) * 10.0)
            intensity = round(0.5733 * (1 - d / 100 * rate / 100), 6)
            trajectory.append({
                "year": year, "scenario_id": sid, "applicable_scope": None,
                "scenario_realization_rate_pct": rate,
                "unit_copper_intensity": {"value": intensity, "unit": "t Cu/MWp_DC"},
                "functional_unit": "t Cu/MWp_DC",
                "delta_vs_baseline": {"value": round(intensity - 0.5733, 6), "unit": "t Cu/MWp_DC"},
                "applied_measure_ids": [],
            })
    return {
        "product_id": "pv_station", "region_id": "China", "region_name": "中国",
        "based_on_research_version": 1, "analysis_run_id": "run-test",
        "functional_unit_candidates": ["t Cu/MWp_DC"],
        "functional_unit_rationale": "r", "chosen_functional_unit": "t Cu/MWp_DC",
        "baseline_year": 2025,
        "baseline_unit_intensity": {"value": 0.5733, "unit": "t Cu/MWp_DC"},
        "baseline_calc_method": "calculate.compute_model_baselines(test)",
        "reduction_measures": [{
            "measure_id": "R01", "measure_name": "焊带减薄", "target_component_ids": ["cu1"],
            "mechanism": "m", "mechanism_category": "structural_optimization",
            "expected_reduction": {"value": 6.03, "unit": "%"}, "maturity": "commercial",
            "earliest_feasible_year": 2025,
            "engineering_case": {"project_or_product_name": "p", "implementing_entity": "e",
                                 "year": 2024, "outcome_description": "o",
                                 "scale": "commercial_deployment", "citations": [CITE]},
            "additional_citations": [CITE],
        }],
        "scenarios": scenarios, "trajectory": trajectory, "model_baselines": baselines,
        "amendment_requests": [], "validation_flags": [], "objections": [], "debate_log": [],
    }


def _run(graph_type, result, research=None):
    return {
        "run_id": "testrun0001", "graph_type": graph_type,
        "product_id": "pv_station", "product_name": "光伏电站",
        "region": "中国", "baseline_year": 2025, "constraints": "",
        "llm_mode": "demo", "research_output": research, "result": result,
    }


def _read(xlsx_bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def _col1_values(ws):
    return [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]


def test_template_structure_preserved():
    tpl = openpyxl.load_workbook(TEMPLATE_PATH)["Template"]
    tpl_headers = {r: [c.value for c in row] for r, row in enumerate(tpl.iter_rows(), 1)
                   if row[0].value in ("分类维度名称", "型号/取值名称", "措施ID", "情景代码", "年份")}

    wb = _read(export_run_to_xlsx(_run("reduction", _reduction_result(), _research_result())))
    assert wb.sheetnames == ["Template"]
    ws = wb["Template"]
    col1 = _col1_values(ws)
    for title in ("1. 研究对象、边界与口径", "2. 主流产品分类与代表型号",
                  "3. 含铜部位清单与单位铜强度证据", "4. 减量化措施清单",
                  "5. 减量化情景定义", "6. 单位铜强度变化路径（基准年–2035）"):
        assert any(str(v or "").startswith(title[:12]) for v in col1), title

    # 列头逐字一致
    out_headers = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v in ("分类维度名称", "型号/取值名称", "措施ID", "情景代码", "年份"):
            out_headers[v] = [ws.cell(row=r, column=c).value for c in range(1, 14)]
    for key, tpl_row in tpl_headers.items():
        marker = tpl_row[0]
        expect = [x for x in tpl_row if x is not None]
        got = [x for x in out_headers[marker] if x is not None]
        assert got == expect, marker

    # 示例行已删除
    assert not any(str(v or "").startswith("❌反面示例") for v in col1)
    assert not any(str(v or "").startswith("↑以上") for v in col1)
    for row in ws.iter_rows():
        for c in row:
            rgb = getattr(c.fill.fgColor, "rgb", None)
            assert not (isinstance(rgb, str) and rgb.endswith("FCE4D6"))


def test_reduction_run_contains_section2_and_3():
    wb = _read(export_run_to_xlsx(_run("reduction", _reduction_result(), _research_result())))
    ws = wb.active
    col1 = _col1_values(ws)
    i2 = col1.index("分类维度名称") + 1  # 列头行号（1-based）
    assert ws.cell(row=i2 + 1, column=3).value == "TOPCon晶硅组件"
    assert ws.cell(row=i2 + 1, column=6).value == 87.6  # 市场占有率(%)
    assert ws.cell(row=i2 + 2, column=13).value == "区域证据不足，采用全球主流型号替代"
    i3 = col1.index("型号/取值名称") + 1
    assert ws.cell(row=i3 + 1, column=6).value == 0.5733  # 型号基线强度
    # 构成权重行 + 高铜参考行
    vals = [ws.cell(row=r, column=1).value for r in range(i3, i3 + 8)]
    assert any(str(v or "").startswith("高铜技术参考基准") for v in vals)


def test_section6_row_count():
    wb = _read(export_run_to_xlsx(_run("reduction", _reduction_result(), _research_result())))
    ws = wb.active
    col1 = _col1_values(ws)
    i6 = col1.index("年份") + 1
    n = 0
    r = i6 + 1
    while r <= ws.max_row and isinstance(ws.cell(row=r, column=1).value, int):
        n += 1
        r += 1
    assert n == 11 * 4  # (2035-2025+1)年 × 4情景 × 1子类别


def test_research_run_export_sections_1_to_3():
    wb = _read(export_run_to_xlsx(_run("research", _research_result())))
    ws = wb.active
    col1 = _col1_values(ws)
    assert "分类维度名称" in col1 and "型号/取值名称" in col1
    assert "措施ID" in col1  # 节结构保留（无数据行）
    i2 = col1.index("分类维度名称") + 1
    assert ws.cell(row=i2 + 1, column=3).value == "TOPCon晶硅组件"


def test_export_preserves_visual_data_band():
    """实际数据行应保留模板的数据区色带，而不是退化成无填充白底。"""
    wb = _read(export_run_to_xlsx(_run("reduction", _reduction_result(), _research_result())))
    ws = wb.active
    section_headers = {ws.cell(row=r, column=1).value: r
                       for r in range(1, ws.max_row + 1)}
    for marker in ("分类维度名称", "型号/取值名称", "措施ID", "情景代码", "年份"):
        header_row = section_headers[marker]
        data_cell = ws.cell(row=header_row + 1, column=1)
        assert data_cell.fill.fill_type == "solid"
        assert data_cell.fill.fgColor.rgb.endswith("E2F0D9")
