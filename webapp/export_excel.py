"""运行结果 Excel 导出（模板填数版）
==================================================================
设计原则：《铜产品减量化模型_通用输出模板.xlsx》是布局的唯一事实源。
导出时运行时读取模板，逐行复制其结构（标题/说明/节标题/列头/说明行），
删除橙色示例行与反面示例行，然后在各节列头之下填入本次运行的数据行；
Section 6 强度路径由 calculate.py 确定性计算后批量写入（模板原生要求）。

与旧版的区别（修复"缺漏/飘移"）：
  - 不再自建列集：列头逐字来自模板，模板改列导出自动跟随；
  - reduction run 的 Section 2/3 从 run["research_output"]（父调研成果）
    与 result["model_baselines"] 填数，修复整节缺失；
  - 基线强度/功能单位/Δmax/路径全部来自确定性计算产出，无 LLM 算术；
  - 示例行（橙色 FCE4D6）与"↑以上为格式示例"说明行在交付件中删除。
"""

from __future__ import annotations
import io
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = _ROOT / "铜产品减量化模型_通用输出模板.xlsx"

EXAMPLE_RGB_SUFFIX = "FCE4D6"  # 橙=示例数据（build_template.EXAMPLE_FILL）

FONT_NAME = "Arial"
TITLE_FONT = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
SECTION_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
HEADER_FONT = Font(name=FONT_NAME, size=9, bold=True)
NOTE_FONT = Font(name=FONT_NAME, size=9, italic=True, color="595959")
BODY_FONT = Font(name=FONT_NAME, size=9)

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="2E75B6")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
FIXED_FILL = PatternFill("solid", fgColor="F2F2F2")
RESULT_FILL = PatternFill("solid", fgColor="E2F0D9")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DIFFUSION_CN = {"linear": "线性扩散", "s_curve": "S型扩散",
                "step": "阶跃扩散", "other": "其他"}

# 各节列头首列标识（用于在模板中定位列头行）
SECTION_HEADER_MARKERS = {
    2: "分类维度名称",
    3: "型号/取值名称",
    4: "措施ID",
    5: "情景代码",
    6: "年份",
}


def _style(cell, font=BODY_FONT, fill=None):
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER


def _is_example_row(row_cells) -> bool:
    for c in row_cells:
        rgb = getattr(c.fill.fgColor, "rgb", None)
        if isinstance(rgb, str) and rgb.endswith(EXAMPLE_RGB_SUFFIX):
            return True
    return False


def _first_cite(obj: dict, key: str = "citations"):
    cites = obj.get(key) or []
    return cites[0] if cites else {}


def _q(obj) -> str:
    if not isinstance(obj, dict):
        return "－"
    v = obj.get("value")
    if v is None:
        return "－"
    return f"{v} {obj.get('unit', '')}".strip()


# ============================================================
# 各节数据行构造（列序与模板列头严格一致）
# ============================================================

def _section2_rows(research: dict) -> list[list]:
    rows = []
    for dim in research.get("classification_dimensions") or []:
        cite0 = _first_cite(dim, "evidence")
        for v in dim.get("values") or []:
            vc = _first_cite(v)
            share = v.get("market_share") or {}
            rows.append([
                dim.get("dimension_name"), dim.get("dimension_category"),
                v.get("value_name"), v.get("typical_spec_range") or "－",
                "－",  # 主要制造商/代表厂商：schema 未承载，留待后续扩展
                share.get("value") if share.get("value") is not None else "－",
                v.get("market_share_caliber") or "－",
                v.get("market_share_year") or "－",
                vc.get("title") or cite0.get("title") or "－",
                vc.get("url") or cite0.get("url") or "－",
                vc.get("confidence") or cite0.get("confidence") or "－",
                v.get("axis_conformity_justification"),
                v.get("region_basis_note") or (
                    "全球主流型号替代" if v.get("region_basis") == "global_fallback"
                    else "区域特有证据"),
            ])
    return rows


def _section3_rows(result: dict, research: dict) -> list[list]:
    baselines = result.get("model_baselines") or []
    rows = []
    if baselines:
        for b in baselines:
            cite0 = _first_cite(b)
            parts = "＋".join(f"{cw['component_name']}{cw['mass']}"
                               for cw in b.get("component_weights") or [])
            rows.append([
                f"{b.get('value_name')}（研究基准）", b.get("applicable_scope") or "(无子类别)",
                "全部含铜部件", parts or "－",
                (b.get("component_weights") or [{}])[0].get("mass_unit", "kg"),
                b["intensity"]["value"], b["intensity"]["unit"],
                cite0.get("title") or "－", cite0.get("url") or "－", 100,
            ])
            for cw in b.get("component_weights") or []:
                contrib = round(b["intensity"]["value"] * cw["weight_pct"] / 100.0, 6)
                rows.append([
                    b.get("value_name"), b.get("applicable_scope") or "(无子类别)",
                    cw["component_name"], f"{cw['mass']} {cw['mass_unit']}",
                    cw["mass_unit"], contrib, b["intensity"]["unit"],
                    "－", "－", cw["weight_pct"],
                ])
        for b in baselines:
            if b.get("is_reference"):
                rows.append([
                    f"高铜技术参考基准（{b.get('applicable_scope') or '无子类别'}）",
                    b.get("applicable_scope") or "(无子类别)",
                    b.get("value_name"), "－", "－",
                    b["intensity"]["value"], b["intensity"]["unit"],
                    b.get("high_copper_reference_note") or "－", "－", 100,
                ])
    else:
        # 仅调研 run：基线尚未由图B计算，先列含铜部位证据
        for c in research.get("copper_components") or []:
            cite0 = _first_cite(c)
            um = c.get("unit_mass") or {}
            rows.append([
                str(c.get("applies_to_dimension_values") or "全部型号"),
                "(无子类别)", c.get("component_name"),
                f"{um.get('value')} {um.get('unit', '')}".strip(),
                um.get("unit", "kg"), "－", "－",
                cite0.get("title") or "－", cite0.get("url") or "－", "－",
            ])
    return rows


def _section4_rows(result: dict) -> list[list]:
    import calculate
    baselines = result.get("model_baselines") or []
    rows = []
    for m in result.get("reduction_measures") or []:
        case = m.get("engineering_case") or {}
        abs_val, abs_unit = calculate.measure_absolute_reduction(m, baselines)
        rows.append([
            m.get("measure_id"), m.get("mechanism_category"), m.get("measure_name"),
            m.get("mechanism"), "、".join(m.get("target_component_ids") or []),
            m.get("applicable_scope") or "(无子类别)",
            abs_val, abs_unit,
            (m.get("expected_reduction") or {}).get("value", "－"),
            "－",  # 可叠加措施ID：schema 未承载
            f"{case.get('project_or_product_name', '')} / "
            f"{case.get('implementing_entity', '')} / {case.get('year', '')}",
            m.get("trade_offs") or "－",
            m.get("maturity"),
        ])
    return rows


def _section5_rows(result: dict) -> list[list]:
    rows = []
    for s in result.get("scenarios") or []:
        cite0 = _first_cite(s)
        rows.append([
            s.get("scenario_id"), s.get("scenario_name"), s.get("scenario_definition"),
            s.get("applicable_scope") or "(无子类别)",
            "、".join(s.get("included_measure_ids") or []) or "无",
            (s.get("full_implementation_reduction_pct") or {}).get("value", "－"),
            s.get("measure_start_year"), s.get("target_achievement_year"),
            s.get("target_achievement_rate_pct"),
            DIFFUSION_CN.get(s.get("diffusion_method"), s.get("diffusion_method")),
            s.get("scenario_rationale"), s.get("notes") or "－",
            cite0.get("url") or "－",
        ])
    return rows


def _section6_rows(result: dict, baseline_year: int) -> list[list]:
    rows = []
    for p in sorted(result.get("trajectory") or [],
                    key=lambda x: (x.get("scenario_id"), x.get("year"))):
        note = "基准年" if (p.get("year") == baseline_year
                           and p.get("scenario_realization_rate_pct") == 0) else ""
        rows.append([
            p.get("year"), p.get("applicable_scope") or "(无子类别)",
            p.get("scenario_id"), p.get("scenario_realization_rate_pct"),
            (p.get("unit_copper_intensity") or {}).get("value"),
            (p.get("unit_copper_intensity") or {}).get("unit"),
            (p.get("delta_vs_baseline") or {}).get("value"), note,
        ])
    return rows


# ============================================================
# 主导出函数
# ============================================================

def export_run_to_xlsx(run: dict) -> bytes:
    result = run.get("result") or {}
    is_reduction = "reduction_measures" in result
    research = result if "classification_dimensions" in result else (run.get("research_output") or {})

    tpl = openpyxl.load_workbook(TEMPLATE_PATH)
    ws_in = tpl["Template"]
    ncols = ws_in.max_column

    # 模板是样式的唯一事实源。输出行数会随研究结果变化，不能直接复制整行，
    # 但可以把模板中的行样式复制到对应的输出行，确保正文填充色、字体、边框、
    # 对齐方式和行高与模板保持一致。
    style_rows = {
        "title": 1,
        "note": 2,
        "field": 5,
        "section": 4,
        "header": {
            2: 16,
            3: 22,
            4: 29,
            5: 35,
            6: 42,
        },
        "data": {
            2: 17,
            3: 23,
            4: 30,
            5: 36,
            6: 43,
        },
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Template")
    ws.sheet_view.showGridLines = False
    for col in range(1, ncols + 1):
        letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[letter].width = ws_in.column_dimensions[letter].width

    section1_values = _section1_values(run, result, is_reduction)
    section_rows = {
        2: _section2_rows(research),
        3: _section3_rows(result, research),
        4: _section4_rows(result) if is_reduction else [],
        5: _section5_rows(result) if is_reduction else [],
        6: _section6_rows(result, run.get("baseline_year", 2025)) if is_reduction else [],
    }
    extra_section1 = _section1_extra(run, result, is_reduction)

    # 第一遍：把模板行翻译成操作序列（不用 insert_rows——openpyxl 的
    # insert_rows 不移动合并区域，会导致合并区与数据行错位）
    ops: list[tuple] = []
    title_done = False
    for row in ws_in.iter_rows():
        vals = [c.value for c in row]
        if not any(v is not None for v in vals):
            ops.append(("blank",))
            continue
        if _is_example_row(row):
            continue  # 橙色示例行不进入交付件
        first = str(vals[0]) if vals[0] is not None else ""
        if first.startswith("↑以上"):
            continue  # 示例说明行不进入交付件
        if not title_done:
            ops.append(("title", f"{run['product_name']} 铜减量化模型"))
            title_done = True
            continue
        if first.startswith("填写说明"):
            ops.append(("note", f"填写说明：本文件由系统按通用输出模板自动填数；"
                                f"运行ID={run['run_id']}；模式={run.get('llm_mode', '')}"))
            continue
        if first[:2] in ("1.", "2.", "3.", "4.", "5.", "6.") and first[2:3] == " ":
            ops.append(("section", first))
            continue
        if first in section1_values:
            ops.append(("field", first, section1_values[first]))
            continue
        marker_sec = next((sec for sec, mk in SECTION_HEADER_MARKERS.items()
                           if first == mk), None)
        if marker_sec is not None:
            if marker_sec == 2:
                for label, value in extra_section1:
                    ops.append(("field", label, value))
            ops.append(("header", marker_sec, vals))
            data = section_rows.get(marker_sec) or []
            if data:
                ops.append(("data", marker_sec, data))
            continue
        ops.append(("note", first))

    # 第二遍：顺序写出，合并区域在最终行号上创建
    out_row = 0
    for op in ops:
        out_row += 1
        kind = op[0]
        if kind == "blank":
            continue
        if kind == "title":
            _write_merged(ws, out_row, op[1], ncols, ws_in, style_rows["title"], center=True)
        elif kind == "note":
            note_row = style_rows["note"]
            if op[1].startswith("【重要】"):
                note_row = 15
            elif op[1].startswith("高铜技术"):
                note_row = 26
            elif op[1].startswith("本节数据"):
                note_row = 41
            _write_merged(ws, out_row, op[1], ncols, ws_in, note_row)
        elif kind == "section":
            _write_merged(ws, out_row, op[1], ncols, ws_in, style_rows["section"])
        elif kind == "field":
            _copy_row_style(ws, ws_in, out_row, style_rows["field"], ncols)
            lc = ws.cell(row=out_row, column=1, value=op[1])
            ws.merge_cells(start_row=out_row, start_column=2,
                           end_row=out_row, end_column=ncols)
            vc = ws.cell(row=out_row, column=2,
                         value=op[2] if op[2] is not None else "－")
        elif kind == "header":
            _copy_row_style(ws, ws_in, out_row, style_rows["header"][op[1]], ncols)
            for i, v in enumerate(op[2], start=1):
                ws.cell(row=out_row, column=i, value=v)
        elif kind == "data":
            out_row -= 1  # 数据行占多行，下面自行推进
            for row_vals in op[2]:
                out_row += 1
                _copy_row_style(ws, ws_in, out_row,
                                style_rows["data"][op[1]], ncols)
                for j, v in enumerate(row_vals, start=1):
                    ws.cell(row=out_row, column=j,
                            value=v if v is not None else "－")

    # 附表：校验标记与质询记录
    out_row = ws.max_row + 2
    _write_merged(ws, out_row, "附：校验标记与质询记录", ncols, ws_in, style_rows["section"])
    out_row += 1
    for i, h in enumerate(["类型", "级别", "描述", "关联ID", "状态"], start=1):
        _copy_cell_style(ws.cell(row=out_row, column=i), ws_in.cell(row=16, column=i))
        ws.cell(row=out_row, column=i, value=h)
    out_row += 1
    for f in result.get("validation_flags") or []:
        for j, v in enumerate(["校验标记", f.get("severity"), f.get("description"),
                               "、".join(f.get("related_ids") or []),
                               "已解决" if f.get("resolved") else "未解决"], start=1):
            _style(ws.cell(row=out_row, column=j, value=v))
        out_row += 1
    for o in result.get("objections") or []:
        for j, v in enumerate([f"质询({o.get('flag_type')})", o.get("severity"),
                               o.get("detail"), "、".join(o.get("target_field_refs") or []),
                               "已回应" if o.get("addressed") else "未回应"], start=1):
            _style(ws.cell(row=out_row, column=j, value=v))
        out_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _copy_cell_style(dst, src):
    # 不能直接复制 _style：模板与新工作簿各自维护样式表，直接复制会产生
    # 无效的样式索引。逐项复制可跨工作簿保留完整视觉样式。
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)
    dst.number_format = src.number_format


def _copy_row_style(ws, ws_in, out_row, template_row, ncols):
    if ws_in.row_dimensions[template_row].height is not None:
        ws.row_dimensions[out_row].height = ws_in.row_dimensions[template_row].height
    for col in range(1, ncols + 1):
        dst = ws.cell(row=out_row, column=col)
        src = ws_in.cell(row=template_row, column=col)
        _copy_cell_style(dst, src)
        # 模板中的橙色只表示“示例数据”。实际运行结果保留数据区的
        # 颜色层次，但改用绿色结果色，避免把真实结果误标为示例。
        if src.fill.fill_type == "solid":
            rgb = getattr(src.fill.fgColor, "rgb", None)
            if isinstance(rgb, str) and rgb.endswith(EXAMPLE_RGB_SUFFIX):
                dst.fill = copy(RESULT_FILL)


def _write_merged(ws, row, text, ncols, ws_in, template_row, center=False):
    _copy_row_style(ws, ws_in, row, template_row, ncols)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    if center:
        c.alignment = copy(c.alignment)
        c.alignment = Alignment(vertical=c.alignment.vertical,
                                horizontal="center",
                                wrap_text=c.alignment.wrap_text,
                                text_rotation=c.alignment.text_rotation,
                                shrink_to_fit=c.alignment.shrink_to_fit,
                                indent=c.alignment.indent)


def _section1_values(run: dict, result: dict, is_reduction: bool) -> dict:
    base_year = run.get("baseline_year", 2025)
    vals = {
        "产品名称": run["product_name"],
        "研究区域": run.get("region") or "－",
        "情景时间范围": f"{base_year}–2035",
        "单位铜强度口径": result.get("chosen_functional_unit") or "－",
        "产品系统边界/统计口径": (
            f"见Section3：{len((run.get('research_output') or result).get('functional_subsystems') or [])}"
            f"个子系统、{len((run.get('research_output') or result).get('copper_components') or [])}个含铜部位"),
        "情景作用对象": f"当年新建的{run['product_name']}（模板默认口径）",
        "减量化的定义": "措施先按部件铜质量×降低%算绝对减铜量，再除以高铜参考基线强度得到可比减量潜力",
        "关键建模口径": "－",
    }
    if is_reduction:
        bi = result.get("baseline_unit_intensity") or {}
        vals["关键建模口径"] = (
            f"{result.get('baseline_calc_method') or '见calculate.py'}；"
            f"基线强度={bi.get('value')} {bi.get('unit', '')}")
    return vals


def _section1_extra(run: dict, result: dict, is_reduction: bool) -> list:
    assessment = (run.get("research_output") or result or {}).get("evidence_assessment") or {}
    extra = [
        ("补充约束", run.get("constraints") or "无"),
        ("运行模式", run.get("llm_mode") or "－"),
        ("运行ID", run["run_id"]),
        ("区域证据充分度",
         f"{assessment.get('score', '－')}分 / {assessment.get('tier', '未评估')}"
         if assessment else "未评估"),
        ("证据口径", assessment.get("evidence_policy", "region_specific") if assessment else "region_specific"),
        ("兜底说明", assessment.get("summary_note") or "无（区域证据充足或尚未评估）"),
    ]
    if is_reduction:
        extra.insert(2, ("分析运行ID", result.get("analysis_run_id", "")))
        extra.insert(3, ("依据调研版本", f"v{result.get('based_on_research_version', '')}"))
    return extra
